import io

import pymupdf
from docx import Document

from src.features.files.file_extraction_service import extract_file

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_pdf_text_is_extracted():
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Quarterly revenue grew twelve percent")
    text = extract_file(doc.tobytes(), "application/pdf")
    assert "Quarterly revenue grew twelve percent" in text


def test_docx_headings_and_tables_are_extracted():
    document = Document()
    document.add_heading("Plan", level=1)
    document.add_paragraph("Ship the audit fixes.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "owner"
    table.rows[0].cells[1].text = "ana"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_file(buffer.getvalue(), DOCX)
    assert "# Plan" in text and "Ship the audit fixes." in text and "owner | ana" in text


def test_csv_and_markdown_and_text():
    assert extract_file(b"a, b\n\n1 ,2\n", "text/csv") == "a | b\n\n1 | 2"
    assert extract_file(b"# Title\nbody", "text/markdown") == "# Title\nbody"
    assert extract_file("café".encode("utf-8-sig"), "text/plain") == "café"
