import csv
import io
import re
import subprocess
from typing import cast

import pymupdf
import pymupdf4llm
from docx import Document
from src.exceptions import ProcessingLimitError
from src.settings import settings


def decode_text(file: bytes) -> str:
    """Decode uploaded text. Windows exports are often cp1252, so fall back to it."""
    try:
        return file.decode("utf-8-sig")
    except UnicodeDecodeError:
        return file.decode("cp1252", errors="replace")


def extract_text(file: bytes) -> str:
    return decode_text(file)


def extract_csv(file: bytes) -> str:
    content = decode_text(file)
    reader = csv.reader(io.StringIO(content))

    rows = []

    for row in reader:
        cleaned_row = [cell.strip() for cell in row]

        if any(cleaned_row):
            rows.append(" | ".join(cleaned_row))

    return "\n\n".join(rows)


def clean_text(text: str) -> str:
    text = re.sub(r"(?<=\w)\n(?=\w)", "", text)

    text = re.sub(r"[ \t]+", " ", text)

    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def extract_pdf(file: bytes) -> str:
    with pymupdf.open(stream=file, filetype="pdf") as document:
        if document.page_count > settings.MAX_PDF_PAGES:
            raise ProcessingLimitError(
                f"This PDF has more than {settings.MAX_PDF_PAGES} pages, which is too many to process."
            )
        text = cast(
            str,
            pymupdf4llm.to_markdown(
                document,
                write_images=False,
                embed_images=False,
            ),
        )

    return clean_text(text)


def extract_docx(file: bytes) -> str:
    document = Document(io.BytesIO(file))
    text = []

    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text.strip()

        if not paragraph_text:
            continue

        style_name = paragraph.style.name if paragraph.style else None

        if style_name and style_name.startswith("Heading"):
            level = style_name.replace("Heading ", "")
            text.append(f"{'#' * int(level)} {paragraph_text}")
        else:
            text.append(paragraph_text)

    for table in document.tables:
        for row in table.rows:
            text.append(" | ".join(cell.text.strip() for cell in row.cells))

    return "\n\n".join(text)


def extract_doc(file: bytes) -> str:
    result = subprocess.run(
        ["antiword", "-"],
        input=file,
        capture_output=True,
        check=True,
    )

    return decode_text(result.stdout).strip()


EXTRACTION_METHODS = {
    "text/plain": extract_text,
    "text/markdown": extract_text,
    "text/csv": extract_csv,
    "application/pdf": extract_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": extract_docx,
    "application/msword": extract_doc,
}


def extract_file(file: bytes, mime_type: str) -> str:
    extraction_method = EXTRACTION_METHODS.get(mime_type)

    if not extraction_method:
        raise ValueError(f"Unsupported file type: {mime_type}")

    text = extraction_method(file)
    if len(text) > settings.MAX_EXTRACTED_CHARS:
        raise ProcessingLimitError(
            "This file contains too much text to process. Split it into smaller files."
        )
    return text
